"""Local Excel mirror of the dashboard table (Reqs.txt §4-11) — no cloud account.

This is a pure *presentation* layer: it takes the already-computed rows from
`core.compute.build_table` (same values, same colours the web dashboard shows) and
places them on a sheet. No calculation or highlight logic lives here — nothing in
`core/compute.py` is touched.

Layout, one block per trading day, stacked top to bottom in date order:
    col A, B : left blank (user fills in later)
    col C    : date, e.g. "08-Sep"
    col D    : day, e.g. "Wed"
    col E..Q : the table's columns 0..12 (13 columns), same text + same cell colour
    then one blank row before the next day's block starts.

Each day always occupies a fixed 27-row block (+1 blank row) regardless of how many
rows are captured so far, so a day's position never shifts as it fills in during the
day or as later days are added.

Where a day's block goes is decided by reading **this file itself** (see
`_day_start_row`) — never from data.db or any other outside record. That's a
deliberate safety property: if data.db is ever deleted/reset (e.g. following the
"start a fresh day" README step) while an older Pappa.xlsx is kept, the next day still
gets appended after whatever is already in the sheet, instead of being placed at row 1
and overwriting it.
"""
from __future__ import annotations

import os
from datetime import date as _date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from core.schedule_times import SCHEDULE

XLSX_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "Pappa.xlsx")

ROWS_PER_DAY = len(SCHEDULE)       # 27
BLOCK_HEIGHT = ROWS_PER_DAY + 1    # + 1 blank line gap between days
DATE_COL = 3                       # C
DAY_COL = 4                        # D
FIRST_DATA_COL = 5                 # E  (columns 0..12 go in E..Q)
NUM_TABLE_COLS = 13

_NO_FILL = PatternFill(fill_type=None)


def _argb(hexcolor: str) -> str:
    return "FF" + hexcolor.lstrip("#").upper()


def _fmt_date(trade_date: str) -> str:
    return _date.fromisoformat(trade_date).strftime("%d-%b")   # "08-Sep"


def _fmt_day(trade_date: str) -> str:
    return _date.fromisoformat(trade_date).strftime("%a")      # "Wed"


def _load_or_create(path: str) -> Workbook:
    if os.path.exists(path):
        return load_workbook(path)
    wb = Workbook()
    wb.active.title = "Pappa"
    return wb


def _day_start_row(ws, date_str: str) -> int:
    """Where today's block goes, decided purely by reading the sheet itself — never
    from data.db or any other outside record. Scans block by block from the top:
    an empty block -> today goes there (first use); a block already stamped with
    today's date -> reuse it (later capture the same day); anything else -> keep
    looking. This means a day's rows can never be reassigned to a different day just
    because some other file (like data.db) was reset or is out of sync."""
    r = 1
    while True:
        existing = ws.cell(r, DATE_COL).value
        if existing is None or existing == date_str:
            return r
        r += BLOCK_HEIGHT


def export_day(trade_date: str, table_rows: list[dict], path: str | None = None) -> None:
    """Write/refresh one day's block. `table_rows` is core.compute.build_table's
    output for that day (list of {text, colors, ...}), already in row-1..27 order.
    `path` overrides where the file lives (e.g. inside a synced OneDrive/Google Drive
    folder) — defaults to Pappa.xlsx in the project folder."""
    path = path or XLSX_PATH
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = _load_or_create(path)
    ws = wb.active
    date_str, day_str = _fmt_date(trade_date), _fmt_day(trade_date)
    start = _day_start_row(ws, date_str)

    for i in range(ROWS_PER_DAY):
        r = start + i
        row = table_rows[i] if i < len(table_rows) else None
        ws.cell(r, DATE_COL, date_str)
        ws.cell(r, DAY_COL, day_str)
        for col in range(NUM_TABLE_COLS):
            cell = ws.cell(r, FIRST_DATA_COL + col)
            cell.value = row["text"].get(col, "") if row else ""
            color = row["colors"].get(col) if row else None
            cell.fill = (PatternFill(start_color=_argb(color), end_color=_argb(color),
                                     fill_type="solid") if color else _NO_FILL)

    wb.save(path)
